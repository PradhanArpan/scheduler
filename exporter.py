import os
import json
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape, A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def get_base_course_code(code):
    if not code:
        return 'DEFAULT'
    c = str(code).upper().strip()
    c = c.replace('-LAB', '').replace('P-LAB', '').replace('(LAB)', '')
    c = c.replace('-A', '').replace('-B', '').replace('-C', '').replace('-D', '')
    if c.endswith('P') and len(c) > 3 and c[-2].isdigit():
        c = c[:-1]
    return c.replace(' ', '').replace('-', '').replace('_', '')

# Hex colors for subject families (pastel tones)
PALETTE_HEX = [
    'E0F2FE', # Sky
    'DCFCE7', # Emerald
    'FEF3C7', # Amber
    'EDE9FE', # Violet
    'FFE4E6', # Rose
    'EEF2FF', # Indigo
    'CCFBF1', # Teal
    'FFEDD5', # Orange
    'FCE7F3', # Fuchsia
    'ECFCCB', # Lime
]

def get_subject_hex(code):
    if not code:
        return 'FFFFFF'
    base = get_base_course_code(code)
    h = sum(ord(ch) for ch in base)
    return PALETTE_HEX[h % len(PALETTE_HEX)]

def get_subject_reportlab_color(code):
    h = get_subject_hex(code)
    return colors.HexColor(f"#{h}")

# =============================================================================
# 1. EXCEL EXPORT (Multi-sheet: Class TT, Faculty TT, Laboratory)
# =============================================================================
def generate_excel(data, sem_type=None):
    if not data:
        return None
    if sem_type is None:
        sem_type = data.get("config", {}).get("semester", "Even Semester")
    
    is_even = "Even" in str(sem_type)
    target_sem = "Even" if is_even else "Odd"

    wb = openpyxl.Workbook()
    # remove default sheet later

    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )

    f_univ = Font(name='Times New Roman', size=13, bold=True)
    f_dept = Font(name='Times New Roman', size=11, bold=True)
    f_sub = Font(name='Times New Roman', size=11, bold=True)
    f_head = Font(name='Times New Roman', size=10, bold=True)
    f_day = Font(name='Times New Roman', size=10, bold=True)
    f_body = Font(name='Times New Roman', size=9, bold=False)
    f_body_bold = Font(name='Times New Roman', size=9, bold=True)

    al_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    al_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    fill_header = PatternFill(start_color='F1F5F9', fill_type='solid')
    fill_break = PatternFill(start_color='FEF9C3', fill_type='solid')
    fill_closed = PatternFill(start_color='F8FAFC', fill_type='solid')

    days = data.get("config", {}).get("days", ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"])
    period_cols = [
        (0, 4, "09:00 - 10:00"),
        (1, 5, "10:00 - 11:00"),
        (2, 6, "11:00 - 12:00"),
        (3, 7, "12:00 - 13:00"),
        ("break", 8, "Break"),
        (5, 9, "14:00 - 15:00"),
        (6, 10, "15:00 - 16:00")
    ]

    timetable = data.get("timetable", [])

    # ---------------------------------------------------------
    # Sheet 1: Class TT
    # ---------------------------------------------------------
    ws_class = wb.create_sheet(title="Class TT")
    ws_class.views.sheetView[0].showGridLines = True
    
    ws_class.column_dimensions['A'].width = 3.0
    ws_class.column_dimensions['B'].width = 3.0
    ws_class.column_dimensions['C'].width = 15.0
    ws_class.column_dimensions['D'].width = 23.0
    ws_class.column_dimensions['E'].width = 23.0
    ws_class.column_dimensions['F'].width = 23.0
    ws_class.column_dimensions['G'].width = 23.0
    ws_class.column_dimensions['H'].width = 6.0
    ws_class.column_dimensions['I'].width = 23.0
    ws_class.column_dimensions['J'].width = 23.0

    batches = [b for b in data.get("batches", []) if b.get("semester_type") == target_sem]

    curr_row = 2
    for b in batches:
        b_id = b["id"]
        b_name = b["name"]

        # University Header
        ws_class.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
        c = ws_class.cell(row=curr_row, column=3, value="CHRIST (DEEMED TO BE UNIVERSITY)\nSchool of Engineering and Technology")
        c.font = f_univ; c.alignment = al_center
        ws_class.row_dimensions[curr_row].height = 34
        curr_row += 1

        ws_class.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
        c = ws_class.cell(row=curr_row, column=3, value="Department of Civil Engineering")
        c.font = f_dept; c.alignment = al_center
        ws_class.row_dimensions[curr_row].height = 18
        curr_row += 1

        ws_class.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
        c = ws_class.cell(row=curr_row, column=3, value=f"{b_name} ({sem_type})")
        c.font = f_sub; c.alignment = al_center
        ws_class.row_dimensions[curr_row].height = 18
        curr_row += 1

        # Table Header
        head_row = curr_row
        ws_class.row_dimensions[head_row].height = 22
        c = ws_class.cell(row=head_row, column=3, value="DAY/TIME")
        c.font = f_head; c.alignment = al_center; c.fill = fill_header; c.border = thin_border

        for p_idx, col_num, time_str in period_cols:
            c = ws_class.cell(row=head_row, column=col_num, value=time_str)
            c.font = f_head; c.alignment = al_center; c.fill = fill_header; c.border = thin_border
        curr_row += 1

        # Day rows
        for day_idx, day_name in enumerate(days):
            r = curr_row
            ws_class.row_dimensions[r].height = 38
            c_day = ws_class.cell(row=r, column=3, value=day_name)
            c_day.font = f_day; c_day.alignment = al_center; c_day.border = thin_border; c_day.fill = fill_header

            c_brk = ws_class.cell(row=r, column=8, value="Break")
            c_brk.font = Font(name='Times New Roman', size=8, bold=True, color='854D0E')
            c_brk.alignment = al_center; c_brk.fill = fill_break; c_brk.border = thin_border

            day_entries = [e for e in timetable if (e.get("batch_id") == b_id or (e.get("batch_name") and b_name in e.get("batch_name"))) and e.get("day") == day_idx]
            period_entry_map = {e.get("period"): e for e in day_entries}

            handled_periods = set()
            for p_idx, col_num, _ in [(0, 4, None), (1, 5, None), (2, 6, None), (3, 7, None), (5, 9, None), (6, 10, None)]:
                if p_idx in handled_periods:
                    continue

                cell = ws_class.cell(row=r, column=col_num)
                cell.border = thin_border

                if day_idx == 5 and p_idx in [5, 6]:
                    cell.value = "-X-"
                    cell.font = f_body_bold; cell.alignment = al_center; cell.fill = fill_closed
                    handled_periods.add(p_idx)
                    continue

                e = period_entry_map.get(p_idx)
                if not e:
                    cell.value = ""
                    cell.alignment = al_center
                    handled_periods.add(p_idx)
                    continue

                code = e.get("course_code", "")
                t_name = e.get("teacher_code") or e.get("teacher_name", "")
                sb_name = e.get("sub_batch_name")
                is_lab = e.get("is_lab", False)
                r_name = e.get("room_name", "")

                label = f"{code}\n{t_name}" if t_name else code
                if r_name:
                    label += f"\n[{r_name}]"
                if sb_name:
                    label = f"[{sb_name}] {label}"

                cell.value = label
                cell.alignment = al_center
                cell.font = f_body_bold if is_lab else f_body
                cell.fill = PatternFill(start_color=get_subject_hex(code), fill_type='solid')

                next_p = p_idx + 1
                next_col = col_num + 1
                if is_lab and next_p in period_entry_map and (next_p, next_col, None) in [(1, 5, None), (3, 7, None), (6, 10, None)]:
                    next_e = period_entry_map[next_p]
                    if next_e.get("course_code") == code:
                        ws_class.merge_cells(start_row=r, start_column=col_num, end_row=r, end_column=next_col)
                        c_next = ws_class.cell(row=r, column=next_col)
                        c_next.border = thin_border
                        handled_periods.add(next_p)

                handled_periods.add(p_idx)

            curr_row += 1

        curr_row += 2

    # ---------------------------------------------------------
    # Sheet 2: Faculty TT
    # ---------------------------------------------------------
    ws_fac = wb.create_sheet(title="Faculty TT")
    ws_fac.views.sheetView[0].showGridLines = True

    ws_fac.column_dimensions['A'].width = 3.0
    ws_fac.column_dimensions['B'].width = 3.0
    ws_fac.column_dimensions['C'].width = 15.0
    ws_fac.column_dimensions['D'].width = 22.0
    ws_fac.column_dimensions['E'].width = 22.0
    ws_fac.column_dimensions['F'].width = 22.0
    ws_fac.column_dimensions['G'].width = 22.0
    ws_fac.column_dimensions['H'].width = 6.0
    ws_fac.column_dimensions['I'].width = 22.0
    ws_fac.column_dimensions['J'].width = 22.0
    ws_fac.column_dimensions['K'].width = 3.0
    ws_fac.column_dimensions['L'].width = 18.0
    ws_fac.column_dimensions['M'].width = 14.0

    curr_row = 2
    for f in data.get("faculty", []):
        f_id = f["id"]
        f_name = f["name"]
        f_code = f.get("code", "")
        f_desig = f.get("designation", "Assistant Professor")

        ws_fac.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
        c = ws_fac.cell(row=curr_row, column=3, value="CHRIST (DEEMED TO BE UNIVERSITY)\nSchool of Engineering and Technology")
        c.font = f_univ; c.alignment = al_center
        ws_fac.row_dimensions[curr_row].height = 34
        curr_row += 1

        ws_fac.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
        c = ws_fac.cell(row=curr_row, column=3, value="Department of Civil Engineering")
        c.font = f_dept; c.alignment = al_center
        ws_fac.row_dimensions[curr_row].height = 18
        curr_row += 1

        ws_fac.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
        c = ws_fac.cell(row=curr_row, column=3, value=f"{f_name} ({f_code})")
        c.font = f_sub; c.alignment = al_center
        ws_fac.row_dimensions[curr_row].height = 18
        curr_row += 1

        head_row = curr_row
        ws_fac.row_dimensions[head_row].height = 22
        c = ws_fac.cell(row=head_row, column=3, value="DAY/TIME")
        c.font = f_head; c.alignment = al_center; c.fill = fill_header; c.border = thin_border

        for p_idx, col_num, time_str in period_cols:
            c = ws_fac.cell(row=head_row, column=col_num, value=time_str)
            c.font = f_head; c.alignment = al_center; c.fill = fill_header; c.border = thin_border

        c_sum = ws_fac.cell(row=head_row, column=12, value="Summary")
        c_sum.font = f_head; c_sum.alignment = al_center; c_sum.fill = fill_header; c_sum.border = thin_border
        ws_fac.merge_cells(start_row=head_row, start_column=12, end_row=head_row, end_column=13)
        curr_row += 1

        th_hours = 0
        lab_contact = 0
        for s in data.get("unmapped_subjects", []):
            if (s.get("semester_type") or "Even") != target_sem:
                continue
            t_ids = s.get("teacher_ids") or ([s["teacher_id"]] if s.get("teacher_id") else [])
            if f_id in t_ids:
                p_cnt = s.get("weekly_periods", 0)
                if s.get("is_lab") or "LAB" in str(s.get("code")).upper():
                    lab_contact += p_cnt
                else:
                    th_hours += p_cnt

        lab_workload = lab_contact / 2
        tot_workload = th_hours + lab_workload
        lab_wl_str = f"{int(lab_workload)}" if lab_workload == int(lab_workload) else f"{lab_workload}"
        tot_wl_str = f"{int(tot_workload)}" if tot_workload == int(tot_workload) else f"{tot_workload}"

        sum_rows_data = [
            ("Lecture Hours", f"{th_hours}"),
            ("Lab Hours", f"{lab_contact} ({lab_wl_str} load)"),
            ("Active Load", f"{th_hours} + {lab_wl_str}"),
            ("Total Workload", f"{tot_wl_str} hrs"),
            ("Designation", f_desig)
        ]

        for day_idx, day_name in enumerate(days):
            r = curr_row
            ws_fac.row_dimensions[r].height = 38
            c_day = ws_fac.cell(row=r, column=3, value=day_name)
            c_day.font = f_day; c_day.alignment = al_center; c_day.border = thin_border; c_day.fill = fill_header

            c_brk = ws_fac.cell(row=r, column=8, value="Break")
            c_brk.font = Font(name='Times New Roman', size=8, bold=True, color='854D0E')
            c_brk.alignment = al_center; c_brk.fill = fill_break; c_brk.border = thin_border

            day_entries = [e for e in timetable if (e.get("teacher_id") == f_id or f_id in (e.get("teacher_ids") or [])) and e.get("day") == day_idx]
            period_entry_map = {e.get("period"): e for e in day_entries}

            handled_periods = set()
            for p_idx, col_num, _ in [(0, 4, None), (1, 5, None), (2, 6, None), (3, 7, None), (5, 9, None), (6, 10, None)]:
                if p_idx in handled_periods:
                    continue

                cell = ws_fac.cell(row=r, column=col_num)
                cell.border = thin_border

                if day_idx == 5 and p_idx in [5, 6]:
                    cell.value = "-X-"
                    cell.font = f_body_bold; cell.alignment = al_center; cell.fill = fill_closed
                    handled_periods.add(p_idx)
                    continue

                e = period_entry_map.get(p_idx)
                if not e:
                    cell.value = ""
                    cell.alignment = al_center
                    handled_periods.add(p_idx)
                    continue

                b_name_s = e.get("batch_name", "")
                code = e.get("course_code", "")
                is_lab = e.get("is_lab", False)

                r_name = e.get("room_name", "")
                label = f"{b_name_s}\n{code}"
                if r_name:
                    label += f"\n[{r_name}]"
                cell.value = label
                cell.alignment = al_center
                cell.font = f_body_bold if is_lab else f_body
                cell.fill = PatternFill(start_color=get_subject_hex(code), fill_type='solid')

                next_p = p_idx + 1
                next_col = col_num + 1
                if is_lab and next_p in period_entry_map and (next_p, next_col, None) in [(1, 5, None), (3, 7, None), (6, 10, None)]:
                    next_e = period_entry_map[next_p]
                    if next_e.get("course_code") == code:
                        ws_fac.merge_cells(start_row=r, start_column=col_num, end_row=r, end_column=next_col)
                        c_next = ws_fac.cell(row=r, column=next_col)
                        c_next.border = thin_border
                        handled_periods.add(next_p)

                handled_periods.add(p_idx)

            s_idx = day_idx
            if s_idx < len(sum_rows_data):
                lbl, val = sum_rows_data[s_idx]
                cL = ws_fac.cell(row=r, column=12, value=lbl)
                cL.font = f_body_bold if lbl == "Total Workload" else f_body
                cL.alignment = al_left; cL.border = thin_border
                cM = ws_fac.cell(row=r, column=13, value=val)
                cM.font = f_body_bold if lbl == "Total Workload" else f_body
                cM.alignment = al_center; cM.border = thin_border
                if lbl == "Total Workload":
                    cL.fill = PatternFill(start_color='EEF2FF', fill_type='solid')
                    cM.fill = PatternFill(start_color='EEF2FF', fill_type='solid')

            curr_row += 1

        curr_row += 2

    # ---------------------------------------------------------
    # Sheet 3: Laboratory
    # ---------------------------------------------------------
    ws_lab = wb.create_sheet(title="Laboratory")
    ws_lab.views.sheetView[0].showGridLines = True

    ws_lab.column_dimensions['A'].width = 3.0
    ws_lab.column_dimensions['B'].width = 3.0
    ws_lab.column_dimensions['C'].width = 15.0
    ws_lab.column_dimensions['D'].width = 24.0
    ws_lab.column_dimensions['E'].width = 24.0
    ws_lab.column_dimensions['F'].width = 24.0
    ws_lab.column_dimensions['G'].width = 24.0
    ws_lab.column_dimensions['H'].width = 6.0
    ws_lab.column_dimensions['I'].width = 24.0
    ws_lab.column_dimensions['J'].width = 24.0

    curr_row = 2
    ws_lab.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
    c = ws_lab.cell(row=curr_row, column=3, value="CHRIST (DEEMED TO BE UNIVERSITY)\nSchool of Engineering and Technology")
    c.font = f_univ; c.alignment = al_center
    ws_lab.row_dimensions[curr_row].height = 34
    curr_row += 1

    ws_lab.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
    c = ws_lab.cell(row=curr_row, column=3, value="Department of Civil Engineering")
    c.font = f_dept; c.alignment = al_center
    ws_lab.row_dimensions[curr_row].height = 18
    curr_row += 1

    ws_lab.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=10)
    c = ws_lab.cell(row=curr_row, column=3, value=f"Laboratory Master Schedule ({sem_type})")
    c.font = f_sub; c.alignment = al_center
    ws_lab.row_dimensions[curr_row].height = 18
    curr_row += 1

    head_row = curr_row
    ws_lab.row_dimensions[head_row].height = 22
    c = ws_lab.cell(row=head_row, column=3, value="DAY/TIME")
    c.font = f_head; c.alignment = al_center; c.fill = fill_header; c.border = thin_border
    for p_idx, col_num, time_str in period_cols:
        c = ws_lab.cell(row=head_row, column=col_num, value=time_str)
        c.font = f_head; c.alignment = al_center; c.fill = fill_header; c.border = thin_border
    curr_row += 1

    for day_idx, day_name in enumerate(days):
        r = curr_row
        ws_lab.row_dimensions[r].height = 44
        c_day = ws_lab.cell(row=r, column=3, value=day_name)
        c_day.font = f_day; c_day.alignment = al_center; c_day.border = thin_border; c_day.fill = fill_header

        c_brk = ws_lab.cell(row=r, column=8, value="Break")
        c_brk.font = Font(name='Times New Roman', size=8, bold=True, color='854D0E')
        c_brk.alignment = al_center; c_brk.fill = fill_break; c_brk.border = thin_border

        day_labs = [e for e in timetable if e.get("is_lab") and e.get("day") == day_idx]
        p_lab_map = {}
        for e in day_labs:
            p_lab_map.setdefault(e.get("period"), []).append(e)

        handled_periods = set()
        for p_idx, col_num, _ in [(0, 4, None), (1, 5, None), (2, 6, None), (3, 7, None), (5, 9, None), (6, 10, None)]:
            if p_idx in handled_periods:
                continue

            cell = ws_lab.cell(row=r, column=col_num)
            cell.border = thin_border

            if day_idx == 5 and p_idx in [5, 6]:
                cell.value = "-X-"
                cell.font = f_body_bold; cell.alignment = al_center; cell.fill = fill_closed
                handled_periods.add(p_idx)
                continue

            entries = p_lab_map.get(p_idx, [])
            if not entries:
                cell.value = ""
                cell.alignment = al_center
                handled_periods.add(p_idx)
                continue

            lines = []
            for e in entries:
                r_name = e.get("room_name", "")
                r_tag = f" [{r_name}]" if r_name else ""
                lines.append(f"{e.get('course_code')} ({e.get('batch_name')}){r_tag}\n{e.get('teacher_name')}")
            cell.value = "\n---\n".join(lines)
            cell.alignment = al_center
            cell.font = f_body_bold
            cell.fill = PatternFill(start_color=get_subject_hex(entries[0].get("course_code")), fill_type='solid')

            next_p = p_idx + 1
            next_col = col_num + 1
            if next_p in p_lab_map and (next_p, next_col, None) in [(1, 5, None), (3, 7, None), (6, 10, None)]:
                ws_lab.merge_cells(start_row=r, start_column=col_num, end_row=r, end_column=next_col)
                c_next = ws_lab.cell(row=r, column=next_col)
                c_next.border = thin_border
                handled_periods.add(next_p)

            handled_periods.add(p_idx)

        curr_row += 1

    if "Sheet" in wb.sheetnames:
        wb.remove(wb["Sheet"])

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out


# =============================================================================
# 2. PDF EXPORT (Class Timetables - Landscape A4)
# =============================================================================
def generate_class_pdf(data, sem_type=None):
    if not data:
        return None
    if sem_type is None:
        sem_type = data.get("config", {}).get("semester", "Even Semester")
    
    is_even = "Even" in str(sem_type)
    target_sem = "Even" if is_even else "Odd"

    out = BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    p_univ = ParagraphStyle(
        'UnivHeader',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=13,
        leading=16,
        alignment=1, # Center
        textColor=colors.HexColor('#0F172A')
    )
    p_dept = ParagraphStyle(
        'DeptHeader',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )
    p_batch = ParagraphStyle(
        'BatchHeader',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#334155')
    )
    p_th_head = ParagraphStyle(
        'TableHead',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8.5,
        leading=10.5,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )
    p_cell_code = ParagraphStyle(
        'CellCode',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8.5,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )
    p_cell_name = ParagraphStyle(
        'CellName',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=6.5,
        leading=8,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )
    p_cell_fac = ParagraphStyle(
        'CellFac',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=7,
        leading=8.5,
        alignment=1,
        textColor=colors.HexColor('#334155')
    )
    p_cell_closed = ParagraphStyle(
        'CellClosed',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8.5,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#64748B')
    )
    p_cell_break = ParagraphStyle(
        'CellBreak',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8,
        leading=9.5,
        alignment=1,
        textColor=colors.HexColor('#854D0E')
    )
    p_ref_head = ParagraphStyle(
        'RefHead',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8,
        leading=10,
        alignment=0,
        textColor=colors.HexColor('#0F172A')
    )
    p_ref_cell = ParagraphStyle(
        'RefCell',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=7.5,
        leading=9.5,
        alignment=0,
        textColor=colors.HexColor('#1E293B')
    )
    p_ref_bold = ParagraphStyle(
        'RefCellBold',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=7.5,
        leading=9.5,
        alignment=0,
        textColor=colors.HexColor('#0F172A')
    )

    days = data.get("config", {}).get("days", ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"])
    timetable = data.get("timetable", [])
    batches = [b for b in data.get("batches", []) if b.get("semester_type") == target_sem]

    col_widths = [62, 102, 102, 102, 102, 50, 102, 102] # total 724 pt
    elements = []

    for b_idx, b in enumerate(batches):
        b_id = b["id"]
        b_name = b["name"]

        elements.append(Paragraph("CHRIST (DEEMED TO BE UNIVERSITY)", p_univ))
        elements.append(Paragraph("School of Engineering and Technology &bull; Department of Civil Engineering", p_dept))
        elements.append(Paragraph(f"Class Timetable: <strong>{b_name}</strong> &bull; {sem_type} (2026-2027)", p_batch))
        elements.append(Spacer(1, 8))

        # Table data
        header_row = [
            Paragraph("DAY / TIME", p_th_head),
            Paragraph("09:00 - 10:00", p_th_head),
            Paragraph("10:00 - 11:00", p_th_head),
            Paragraph("11:00 - 12:00", p_th_head),
            Paragraph("12:00 - 13:00", p_th_head),
            Paragraph("<b>BREAK</b><br/><font size=6.5>13:00-14:00</font>", p_th_head),
            Paragraph("14:00 - 15:00", p_th_head),
            Paragraph("15:00 - 16:00", p_th_head)
        ]
        table_data = [header_row]

        table_styles = [
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#94A3B8')),
            ('BOX', (0, 0), (-1, -1), 1.2, colors.HexColor('#0F172A')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('LEFTPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2),
        ]

        row_index = 1
        for day_idx, day_name in enumerate(days):
            day_entries = [e for e in timetable if (e.get("batch_id") == b_id or (e.get("batch_name") and b_name in e.get("batch_name"))) and e.get("day") == day_idx]
            p_map = {e.get("period"): e for e in day_entries}

            row_cells = [Paragraph(day_name, p_th_head)]
            table_styles.append(('BACKGROUND', (0, row_index), (0, row_index), colors.HexColor('#F8FAFC')))

            # Periods mapping to table col index:
            # 0 -> 1, 1 -> 2, 2 -> 3, 3 -> 4, break -> 5, 5 -> 6, 6 -> 7
            p_to_col = {0: 1, 1: 2, 2: 3, 3: 4, 5: 6, 6: 7}
            row_cells_temp = {}
            handled_periods = set()

            for p_num, c_col in sorted(p_to_col.items()):
                if p_num in handled_periods:
                    continue

                if day_idx == 5 and p_num in [5, 6]:
                    row_cells_temp[c_col] = Paragraph("-X-", p_cell_closed)
                    table_styles.append(('BACKGROUND', (c_col, row_index), (c_col, row_index), colors.HexColor('#F8FAFC')))
                    handled_periods.add(p_num)
                    continue

                e = p_map.get(p_num)
                if not e:
                    row_cells_temp[c_col] = Paragraph("", p_cell_name)
                    handled_periods.add(p_num)
                    continue

                code = e.get("course_code", "")
                c_name = e.get("course_name", "") or code
                t_name = e.get("teacher_name", "") or e.get("teacher_code", "")
                sb_name = e.get("sub_batch_name")
                is_lab = e.get("is_lab", False)

                code_display = f"<b>{code}</b>"
                if sb_name:
                    code_display += f" [{sb_name}]"

                cell_p = [
                    Paragraph(code_display, p_cell_code),
                    Paragraph(c_name, p_cell_name),
                    Paragraph(t_name, p_cell_fac)
                ]

                row_cells_temp[c_col] = cell_p
                table_styles.append(('BACKGROUND', (c_col, row_index), (c_col, row_index), get_subject_reportlab_color(code)))

                # Check if next period merges
                next_p = p_num + 1
                next_c = c_col + 1
                if is_lab and next_p in p_map and next_p in [1, 3, 6]:
                    if p_map[next_p].get("course_code") == code:
                        table_styles.append(('SPAN', (c_col, row_index), (next_c, row_index)))
                        row_cells_temp[next_c] = Paragraph("", p_cell_name)
                        handled_periods.add(next_p)

                handled_periods.add(p_num)

            # Assemble row in exact column order (0 to 7)
            row_cells.append(row_cells_temp.get(1, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(2, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(3, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(4, Paragraph("", p_cell_name)))
            row_cells.append(Paragraph("<b>BREAK</b>", p_cell_break))
            table_styles.append(('BACKGROUND', (5, row_index), (5, row_index), colors.HexColor('#FEF9C3')))
            row_cells.append(row_cells_temp.get(6, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(7, Paragraph("", p_cell_name)))

            table_data.append(row_cells)
            row_index += 1

        t = Table(table_data, colWidths=col_widths, rowHeights=[22] + [42]*len(days))
        t.setStyle(TableStyle(table_styles))
        elements.append(t)

        # Course Code, Subject Name, Faculty In-Charge Reference Table below timetable
        batch_legend = {}
        batch_entries = [e for e in timetable if (e.get("batch_id") == b_id or (e.get("batch_name") and b_name in e.get("batch_name")))]
        for e in batch_entries:
            code = e.get("course_code")
            if not code:
                continue
            c_name = e.get("course_name") or code
            t_name = e.get("teacher_name") or e.get("teacher_code") or ""
            sb_name = e.get("sub_batch_name")
            key = (code, c_name)
            if key not in batch_legend:
                batch_legend[key] = set()
            if t_name:
                if sb_name:
                    batch_legend[key].add(f"{t_name} ({sb_name})")
                else:
                    batch_legend[key].add(t_name)

        if batch_legend:
            items = sorted(batch_legend.items())
            if len(items) <= 6:
                ref_data = [
                    [
                        Paragraph("Course Code", p_ref_head),
                        Paragraph("Subject Name", p_ref_head),
                        Paragraph("Faculty In-Charge", p_ref_head)
                    ]
                ]
                for (c_code, c_name), fac_set in items:
                    fac_str = ", ".join(sorted(fac_set)) if fac_set else "—"
                    ref_data.append([
                        Paragraph(c_code, p_ref_bold),
                        Paragraph(c_name, p_ref_cell),
                        Paragraph(fac_str, p_ref_cell)
                    ])

                ref_table = Table(ref_data, colWidths=[90, 384, 250])
            else:
                # 2-column side-by-side layout for large course sets to prevent multi-page spill
                half = (len(items) + 1) // 2
                left_items = items[:half]
                right_items = items[half:]
                ref_data = [
                    [
                        Paragraph("Code", p_ref_head),
                        Paragraph("Subject Name", p_ref_head),
                        Paragraph("Faculty", p_ref_head),
                        Paragraph("Code", p_ref_head),
                        Paragraph("Subject Name", p_ref_head),
                        Paragraph("Faculty", p_ref_head)
                    ]
                ]
                for i in range(half):
                    (l_code, l_name), l_fac = left_items[i]
                    l_str = ", ".join(sorted(l_fac)) if l_fac else "—"
                    if i < len(right_items):
                        (r_code, r_name), r_fac = right_items[i]
                        r_str = ", ".join(sorted(r_fac)) if r_fac else "—"
                        ref_data.append([
                            Paragraph(l_code, p_ref_bold),
                            Paragraph(l_name, p_ref_cell),
                            Paragraph(l_str, p_ref_cell),
                            Paragraph(r_code, p_ref_bold),
                            Paragraph(r_name, p_ref_cell),
                            Paragraph(r_str, p_ref_cell)
                        ])
                    else:
                        ref_data.append([
                            Paragraph(l_code, p_ref_bold),
                            Paragraph(l_name, p_ref_cell),
                            Paragraph(l_str, p_ref_cell),
                            Paragraph("", p_ref_cell),
                            Paragraph("", p_ref_cell),
                            Paragraph("", p_ref_cell)
                        ])
                ref_table = Table(ref_data, colWidths=[65, 177, 120, 65, 177, 120])

            ref_table.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('BOX', (0, 0), (-1, -1), 0.8, colors.HexColor('#94A3B8')),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 1.8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 1.8),
                ('LEFTPADDING', (0, 0), (-1, -1), 3),
                ('RIGHTPADDING', (0, 0), (-1, -1), 3),
            ]))

            elements.append(Spacer(1, 6))
            elements.append(Paragraph("<b>Course &amp; Faculty Reference List:</b>", ParagraphStyle('RefTitle', fontName='Times-Bold', fontSize=8, leading=10, textColor=colors.HexColor('#1E293B'))))
            elements.append(Spacer(1, 2))
            elements.append(ref_table)

        if b_idx < len(batches) - 1:
            elements.append(PageBreak())

    doc.build(elements)
    out.seek(0)
    return out


# =============================================================================
# 3. PDF EXPORT (Faculty Timetables - Landscape A4)
# =============================================================================
def generate_faculty_pdf(data, sem_type=None):
    if not data:
        return None
    if sem_type is None:
        sem_type = data.get("config", {}).get("semester", "Even Semester")
    
    is_even = "Even" in str(sem_type)
    target_sem = "Even" if is_even else "Odd"

    out = BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    p_univ = ParagraphStyle(
        'UnivHeader',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=13,
        leading=16,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )
    p_dept = ParagraphStyle(
        'DeptHeader',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )
    p_fac = ParagraphStyle(
        'FacHeader',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=11,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#334155')
    )
    p_th_head = ParagraphStyle(
        'TableHead',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8.5,
        leading=10.5,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )
    p_cell_code = ParagraphStyle(
        'CellCodeFac',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8.5,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )
    p_cell_name = ParagraphStyle(
        'CellNameFac',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=6.5,
        leading=8,
        alignment=1,
        textColor=colors.HexColor('#1E293B')
    )
    p_cell_batch = ParagraphStyle(
        'CellBatchFac',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=7,
        leading=8.5,
        alignment=1,
        textColor=colors.HexColor('#334155')
    )
    p_cell_closed = ParagraphStyle(
        'CellClosed',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8.5,
        leading=10,
        alignment=1,
        textColor=colors.HexColor('#64748B')
    )
    p_cell_break = ParagraphStyle(
        'CellBreakFac',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8,
        leading=9.5,
        alignment=1,
        textColor=colors.HexColor('#854D0E')
    )
    p_ref_head = ParagraphStyle(
        'RefHeadFac',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=8,
        leading=10,
        alignment=0,
        textColor=colors.HexColor('#0F172A')
    )
    p_ref_cell = ParagraphStyle(
        'RefCellFac',
        parent=styles['Normal'],
        fontName='Times-Roman',
        fontSize=7.5,
        leading=9.5,
        alignment=0,
        textColor=colors.HexColor('#1E293B')
    )
    p_ref_bold = ParagraphStyle(
        'RefCellBoldFac',
        parent=styles['Normal'],
        fontName='Times-Bold',
        fontSize=7.5,
        leading=9.5,
        alignment=0,
        textColor=colors.HexColor('#0F172A')
    )

    days = data.get("config", {}).get("days", ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"])
    timetable = data.get("timetable", [])
    faculty = data.get("faculty", [])

    col_widths = [60, 84, 84, 84, 84, 48, 84, 84, 120] # total 732 pt
    elements = []

    for f_idx, f in enumerate(faculty):
        f_id = f["id"]
        f_name = f["name"]
        f_code = f.get("code", "")
        f_desig = f.get("designation", "Assistant Professor")

        # Calculate Workload (Theory + Lab / 2)
        th_hours = 0
        lab_contact = 0
        for s in data.get("unmapped_subjects", []):
            if (s.get("semester_type") or "Even") != target_sem:
                continue
            t_ids = s.get("teacher_ids") or ([s["teacher_id"]] if s.get("teacher_id") else [])
            if f_id in t_ids:
                p_cnt = s.get("weekly_periods", 0)
                if s.get("is_lab") or "LAB" in str(s.get("code")).upper():
                    lab_contact += p_cnt
                else:
                    th_hours += p_cnt

        lab_workload = lab_contact / 2
        tot_workload = th_hours + lab_workload
        lab_wl_str = f"{int(lab_workload)}" if lab_workload == int(lab_workload) else f"{lab_workload}"
        tot_wl_str = f"{int(tot_workload)}" if tot_workload == int(tot_workload) else f"{tot_workload}"

        elements.append(Paragraph("CHRIST (DEEMED TO BE UNIVERSITY)", p_univ))
        elements.append(Paragraph("School of Engineering and Technology &bull; Department of Civil Engineering", p_dept))
        elements.append(Paragraph(f"Faculty Timetable: <strong>{f_name} ({f_code})</strong> &bull; {f_desig}", p_fac))
        elements.append(Spacer(1, 8))

        header_row = [
            Paragraph("DAY / TIME", p_th_head),
            Paragraph("09:00 - 10:00", p_th_head),
            Paragraph("10:00 - 11:00", p_th_head),
            Paragraph("11:00 - 12:00", p_th_head),
            Paragraph("12:00 - 13:00", p_th_head),
            Paragraph("<b>BREAK</b><br/><font size=6.5>13:00-14:00</font>", p_th_head),
            Paragraph("14:00 - 15:00", p_th_head),
            Paragraph("15:00 - 16:00", p_th_head),
            Paragraph("Workload Summary", p_th_head)
        ]
        table_data = [header_row]

        table_styles = [
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#94A3B8')),
            ('BOX', (0, 0), (-1, -1), 1.2, colors.HexColor('#0F172A')),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('LEFTPADDING', (0, 0), (-1, -1), 2),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2),
        ]

        summary_lines = [
            f"<b>Lecture Hours:</b> {th_hours} hrs",
            f"<b>Lab Hours:</b> {lab_contact}h ({lab_wl_str} load)",
            f"<b>Active Load:</b> {th_hours} + {lab_wl_str}",
            f"<b>Total Workload:</b> <font color='#4338CA'><b>{tot_wl_str} hrs</b></font>",
            f"<b>Designation:</b> {f_desig}",
            f"<b>Department:</b> Civil Engg"
        ]

        row_index = 1
        for day_idx, day_name in enumerate(days):
            day_entries = [e for e in timetable if (e.get("teacher_id") == f_id or f_id in (e.get("teacher_ids") or [])) and e.get("day") == day_idx]
            p_map = {e.get("period"): e for e in day_entries}

            row_cells = [Paragraph(day_name, p_th_head)]
            table_styles.append(('BACKGROUND', (0, row_index), (0, row_index), colors.HexColor('#F8FAFC')))

            p_to_col = {0: 1, 1: 2, 2: 3, 3: 4, 5: 6, 6: 7}
            row_cells_temp = {}
            handled_periods = set()

            for p_num, c_col in sorted(p_to_col.items()):
                if p_num in handled_periods:
                    continue

                if day_idx == 5 and p_num in [5, 6]:
                    row_cells_temp[c_col] = Paragraph("-X-", p_cell_closed)
                    table_styles.append(('BACKGROUND', (c_col, row_index), (c_col, row_index), colors.HexColor('#F8FAFC')))
                    handled_periods.add(p_num)
                    continue

                e = p_map.get(p_num)
                if not e:
                    row_cells_temp[c_col] = Paragraph("", p_cell_name)
                    handled_periods.add(p_num)
                    continue

                b_name_s = e.get("batch_name", "")
                code = e.get("course_code", "")
                c_name = e.get("course_name", "") or code
                sb_name = e.get("sub_batch_name")
                is_lab = e.get("is_lab", False)

                batch_display = b_name_s
                if sb_name and sb_name not in batch_display:
                    batch_display += f" [{sb_name}]"

                cell_p = [
                    Paragraph(f"<b>{code}</b>", p_cell_code),
                    Paragraph(c_name, p_cell_name),
                    Paragraph(batch_display, p_cell_batch)
                ]

                row_cells_temp[c_col] = cell_p
                table_styles.append(('BACKGROUND', (c_col, row_index), (c_col, row_index), get_subject_reportlab_color(code)))

                next_p = p_num + 1
                next_c = c_col + 1
                if is_lab and next_p in p_map and next_p in [1, 3, 6]:
                    if p_map[next_p].get("course_code") == code:
                        table_styles.append(('SPAN', (c_col, row_index), (next_c, row_index)))
                        row_cells_temp[next_c] = Paragraph("", p_cell_name)
                        handled_periods.add(next_p)

                handled_periods.add(p_num)

            row_cells.append(row_cells_temp.get(1, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(2, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(3, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(4, Paragraph("", p_cell_name)))
            row_cells.append(Paragraph("<b>BREAK</b>", p_cell_break))
            table_styles.append(('BACKGROUND', (5, row_index), (5, row_index), colors.HexColor('#FEF9C3')))
            row_cells.append(row_cells_temp.get(6, Paragraph("", p_cell_name)))
            row_cells.append(row_cells_temp.get(7, Paragraph("", p_cell_name)))

            # Summary item on col 8
            s_text = summary_lines[day_idx] if day_idx < len(summary_lines) else ""
            row_cells.append(Paragraph(s_text, ParagraphStyle('SumItem', fontName='Times-Roman', fontSize=8, leading=10.5, alignment=0, textColor=colors.HexColor('#0F172A'))))
            table_styles.append(('BACKGROUND', (8, row_index), (8, row_index), colors.HexColor('#F8FAFC')))

            table_data.append(row_cells)
            row_index += 1

        t = Table(table_data, colWidths=col_widths, rowHeights=[22] + [42]*len(days))
        t.setStyle(TableStyle(table_styles))
        elements.append(t)

        # Class / Batch, Subject Code, Subject Name Reference Table below timetable
        fac_legend = {}
        fac_entries = [e for e in timetable if (e.get("teacher_id") == f_id or f_id in (e.get("teacher_ids") or []))]
        for e in fac_entries:
            code = e.get("course_code")
            if not code:
                continue
            c_name = e.get("course_name") or code
            b_name_s = e.get("batch_name", "")
            sb_name = e.get("sub_batch_name")
            if sb_name and sb_name not in b_name_s:
                b_name_s += f" ({sb_name})"
            key = (b_name_s, code, c_name)
            fac_legend[key] = True

        if fac_legend:
            ref_data = [
                [
                    Paragraph("Class / Batch", p_ref_head),
                    Paragraph("Subject Code", p_ref_head),
                    Paragraph("Subject Name", p_ref_head)
                ]
            ]
            for (b_name_s, c_code, c_name) in sorted(fac_legend.keys()):
                ref_data.append([
                    Paragraph(b_name_s, p_ref_bold),
                    Paragraph(c_code, p_ref_bold),
                    Paragraph(c_name, p_ref_cell)
                ])

            ref_table = Table(ref_data, colWidths=[150, 110, 472])
            ref_table.setStyle(TableStyle([
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ('BOX', (0, 0), (-1, -1), 0.8, colors.HexColor('#94A3B8')),
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('TOPPADDING', (0, 0), (-1, -1), 2.5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ]))

            elements.append(Spacer(1, 8))
            elements.append(Paragraph("<b>Teaching Assignments &amp; Subject Reference:</b>", ParagraphStyle('RefTitleFac', fontName='Times-Bold', fontSize=8, leading=10, textColor=colors.HexColor('#1E293B'))))
            elements.append(Spacer(1, 3))
            elements.append(ref_table)

        if f_idx < len(faculty) - 1:
            elements.append(PageBreak())

    doc.build(elements)
    out.seek(0)
    return out
